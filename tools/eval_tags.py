"""정답 태그와 자동 태그 비교.

1단계 (두 사람 표시 비교):
  python tools\\eval_tags.py --a samples\\labeling\\label_A.xlsx --b samples\\labeling\\label_B.xlsx
  -> 태그별 일치율, 서로 다른 칸 목록, samples/labeling/gold_draft.csv (일치한 칸만 채움)
  gold_draft.csv 의 빈칸을 두 사람이 의논해 채워 gold.csv 로 저장한다 (값: 1 / 0 / ?).

2단계 (자동 태그 평가):
  python tools\\eval_tags.py --gold samples\\labeling\\gold.csv
  -> 태그별 잘못 붙인 것(오탐)과 놓친 것(누락), 상품별 오류 목록.  '?' 칸은 제외한다.
  자동 태그는 현재 matching.py 로 새로 계산한다 (이름 + 판매자 원문 + 이미지에서 읽은 글).
이 평가는 같은 50개에서 뽑은 개발용 세트라서 독립 검증이 아니다.
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import matching  # noqa: E402

LAB = ROOT / "samples" / "labeling"


def read_sheet(path: str) -> tuple[dict[str, dict[str, str]], dict[str, str], list[str]]:
    ws = load_workbook(path, data_only=True)["표시"]
    head = [str(c.value or "") for c in ws[1]]
    tag_cols = {h.split("\n")[0]: i for i, h in enumerate(head) if ":" in h.split("\n")[0]}
    labels, names, empty = {}, {}, []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[1]:
            continue
        pid = str(row[1])
        names[pid] = str(row[4] or "")
        labels[pid] = {}
        for t, i in tag_cols.items():
            v = row[i]
            v = "" if v is None else str(v).strip()
            if v not in ("1", "0", "?"):
                empty.append(f"{pid} {t}")
                v = ""
            labels[pid][t] = v
    return labels, names, empty


def compare_raters(a_path: str, b_path: str) -> None:
    A, names, ea = read_sheet(a_path)
    B, _, eb = read_sheet(b_path)
    if ea or eb:
        print(f"빈칸 또는 잘못된 값: A {len(ea)}칸, B {len(eb)}칸 (처음 5개: {(ea + eb)[:5]}) -> 채우고 다시 실행하세요.")
    tags = sorted({t for d in A.values() for t in d})
    ids = [i for i in A if i in B]
    print(f"\n상품 {len(ids)}개 x 태그 {len(tags)}개 = {len(ids) * len(tags)}칸")
    tot = same = 0
    diff_rows, per_tag = [], {}
    gold = {}
    for i in ids:
        gold[i] = {}
        for t in tags:
            a, b = A[i][t], B[i][t]
            tot += 1
            if a and a == b:
                same += 1
                gold[i][t] = a
            else:
                gold[i][t] = ""
                diff_rows.append((i, t, a, b))
            per_tag.setdefault(t, [0, 0])
            per_tag[t][0] += 1
            per_tag[t][1] += int(bool(a) and a == b)
    print(f"전체 일치 {same}/{tot} ({same / tot:.1%})")
    print("\n[태그별 일치]  (둘 다 1이거나 둘 다 0이어도 일치로 센다. 대부분 0이라 일치율이 높게 나오는 점에 주의)")
    for t, (n, s) in sorted(per_tag.items(), key=lambda x: x[1][1] / x[1][0]):
        if s < n:
            print(f"  {t:<22} {s}/{n}")
    print(f"\n[서로 다른 칸 {len(diff_rows)}개]  -> 의논해서 gold_draft.csv 의 빈칸을 채우세요")
    for i, t, a, b in diff_rows[:60]:
        print(f"  {i:>9} {t:<22} A={a or '-'} B={b or '-'} | {names[i][:28]}")
    out = LAB / "gold_draft.csv"
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["id", "name"] + tags)
        for i in ids:
            w.writerow([i, names[i]] + [gold[i][t] for t in tags])
    print(f"\n저장: {out}")


def evaluate(gold_path: str) -> None:
    rules = matching.load_rules()
    items = {x["id"]: x for x in json.loads((ROOT / "samples" / "real_50.json").read_text(encoding="utf-8"))["items"]}
    vis = {x["id"]: x.get("text", "") for x in json.loads((ROOT / "samples" / "real_50_vision.json").read_text(encoding="utf-8"))["items"]}
    rows = list(csv.DictReader(open(gold_path, encoding="utf-8-sig")))
    tags = [k for k in rows[0] if ":" in k]
    stat = {t: {"TP": 0, "FP": 0, "FN": 0, "TN": 0, "제외": 0} for t in tags}
    errors = []
    unfilled = 0
    for r in rows:
        it = items[r["id"]]
        auto = set(matching.tag_item_full(it["name"], (it.get("text") or "") + "\n" + (vis.get(r["id"]) or ""), rules)["tags"])
        for t in tags:
            g = (r[t] or "").strip()
            if g == "":
                unfilled += 1
                continue
            if g == "?":
                stat[t]["제외"] += 1
                continue
            has = t in auto
            if g == "1" and has:
                stat[t]["TP"] += 1
            elif g == "1" and not has:
                stat[t]["FN"] += 1
                errors.append((r["id"], t, "놓침", it["name"]))
            elif g == "0" and has:
                stat[t]["FP"] += 1
                errors.append((r["id"], t, "잘못 붙임", it["name"]))
            else:
                stat[t]["TN"] += 1
    if unfilled:
        print(f"주의: 합의되지 않은 빈칸 {unfilled}칸은 제외했습니다.")
    print(f"\n상품 {len(rows)}개. 자동 태그는 현재 matching.py 기준.")
    print(f"{'태그':<22}{'정답1':>6}{'맞게 붙임':>9}{'잘못 붙임':>9}{'놓침':>6}")
    TP = FP = FN = 0
    for t in tags:
        s = stat[t]
        pos = s["TP"] + s["FN"]
        if pos + s["FP"] == 0:
            continue
        TP, FP, FN = TP + s["TP"], FP + s["FP"], FN + s["FN"]
        print(f"{t:<22}{pos:>6}{s['TP']:>9}{s['FP']:>9}{s['FN']:>6}")
    prec = TP / (TP + FP) if TP + FP else float("nan")
    rec = TP / (TP + FN) if TP + FN else float("nan")
    print(f"\n전체  맞게 붙임 {TP} / 잘못 붙임 {FP} / 놓침 {FN}   정밀도 {prec:.1%}  재현율 {rec:.1%}")
    print("(잘못 붙임이 많으면 규칙이 과하게 잡는 것, 놓침이 많으면 규칙이 못 읽거나 상품에 정보가 없는 것)")
    print("\n[오류 목록]")
    for pid, t, kind, name in sorted(errors, key=lambda x: (x[2], x[1])):
        print(f"  {kind:<6} {t:<22} {pid:>9} | {name[:30]}")
    out = LAB / "eval_result.json"
    out.write_text(json.dumps({"stat": stat, "errors": errors}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\n저장: {out}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--a")
    ap.add_argument("--b")
    ap.add_argument("--gold")
    args = ap.parse_args()
    if args.gold:
        evaluate(args.gold)
    elif args.a and args.b:
        compare_raters(args.a, args.b)
    else:
        ap.error("--a 와 --b (1단계) 또는 --gold (2단계)를 주세요")


if __name__ == "__main__":
    main()
