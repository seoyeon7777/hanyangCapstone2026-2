"""정답 태그 표시용 엑셀 만들기.   python tools\\make_label_sheets.py [--n 24] [--seed 20261012]

real_50.json 에서 출처×카테고리로 나눠 무작위 추출(고정 seed)한 상품을 두 사람이 독립적으로 표시할 수 있게
label_A.xlsx / label_B.xlsx 두 파일로 만든다. 자동 태그는 보여 주지 않는다 (먼저 보면 따라가게 된다).
이 50개에서 뽑으므로 '개발용 세트' 평가이다. 독립 검증이라고 부르지 않는다.
출력: samples/labeling/label_A.xlsx, label_B.xlsx, items_manifest.json
"""
from __future__ import annotations

import argparse
import json
import random
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import matching  # noqa: E402

TAG_KO = {
    "fit:over": "오버핏·루즈핏·오버사이즈", "fit:slim": "슬림핏", "len:crop": "크롭", "len:long": "롱(맥시·롱스커트·롱원피스·롱코트)",
    "sil:wide": "와이드", "sil:A": "A라인·플레어", "sil:IH": "일자·스트레이트", "det:ribbon": "리본", "det:ruffle": "러플·프릴",
    "item:tailored_jacket": "테일러드 재킷", "item:tight_skirt": "타이트 스커트",
    "mat:cotton": "면", "mat:wool": "울", "mat:nylon": "나일론", "mat:silk": "실크", "mat:linen": "린넨", "mat:cashmere": "캐시미어",
    "mat:denim": "데님", "mat:knit": "니트", "mat:leather": "가죽", "mat:corduroy": "코듀로이", "mat:chiffon": "시폰", "mat:velvet": "벨벳",
    "mat:fur": "퍼", "mat:mohair": "모헤어", "mat:mustang": "무스탕",
}
FONT = "Arial"
GUIDE = [
    ("목적", "옷의 속성을 제대로 읽었는지 확인하는 정답지입니다. 어떤 체형에 추천할지는 판단하지 않습니다."),
    ("독립 표시", "다른 사람의 표시나 자동 태그를 보지 말고 혼자 표시하세요. 끝나면 합의 때 서로 다른 칸만 같이 봅니다."),
    ("셀 값", "노란 칸에 1(있음) / 0(없음) / ?(자료로 판단 불가)를 고릅니다. 빈칸은 안 됩니다."),
    ("근거", "상품명·판매자 원문·이미지에 실제로 쓰여 있거나 보이는 것만 표시합니다. 일반 상식이나 브랜드 이미지로 추측하지 않습니다."),
    ("소재 30% 기준", "혼용률이 있으면 겉감 기준으로 30% 이상인 섬유만 1입니다. 안감·배색·장식 부분의 혼용률은 제외합니다."),
    ("겉감 우선", "겉감과 안감/배색 혼용률이 따로 있으면 겉감만 봅니다. 겉감 표시가 없고 혼용률이 하나뿐이면 그것을 겉감으로 봅니다."),
    ("혼용률 없을 때", "설명에 '울 부클 소재', '데님' 같이 소재가 직접 쓰여 있으면 1, 없으면 0입니다. 니트·데님·가죽·코듀로이처럼 섬유가 아닌 소재는 혼용률과 무관하게 소재 자체로 봅니다."),
    ("핏 태그", "판매자가 오버핏·루즈핏·오버사이즈, 슬림핏이라고 쓴 경우만 1입니다. 사진만 보고 판단하지 않습니다. '슬림핏 적당함 루즈핏' 같은 선택지 목록과 세탁 안내표의 소재 제목(LEATHER, DENIM 등)은 상품 정보가 아닙니다."),
    ("일자·스트레이트", "일자·스트레이트 실루엣만 1입니다. H라인을 같은 태그로 볼지는 자문 후 정해지므로, H라인이라고만 쓰인 상품은 ?로 두고 메모에 'H라인'이라고 적어 주세요."),
    ("메모", "판단이 어려웠던 이유를 한 줄 적어 두면 합의 때 빨라집니다."),
    ("예시 (표시 방법)", "예: 설명에 '와이드 팬츠, 면 100%'만 있으면 와이드=1, 면=1, 나머지 태그=0."),
]


def pick(items: list[dict], n: int, seed: int) -> list[dict]:
    groups: dict[tuple, list[dict]] = defaultdict(list)
    for it in items:
        groups[(it["src"], it["cat"])].append(it)
    rng = random.Random(seed)
    keys = sorted(groups)
    quota = {k: 0 for k in keys}
    # 칸마다 1개씩 돌려가며 n개가 찰 때까지 배정 (한 칸의 상품 수를 넘지 않게)
    while sum(quota.values()) < n:
        moved = False
        for k in keys:
            if sum(quota.values()) >= n:
                break
            if quota[k] < len(groups[k]):
                quota[k] += 1
                moved = True
        if not moved:
            break
    out = []
    for k in keys:
        out += rng.sample(groups[k], quota[k])
    return sorted(out, key=lambda x: x["id"])


def build(path: Path, rows: list[dict], tags: list[str], who: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "안내"
    ws["A1"] = f"정답 태그 표시 ({who})"
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    for i, (k, v) in enumerate(GUIDE, start=3):
        ws.cell(i, 1, k).font = Font(name=FONT, bold=True)
        c = ws.cell(i, 2, v)
        c.font = Font(name=FONT)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 110
    r0 = 3 + len(GUIDE) + 1
    ws.cell(r0, 1, "태그 목록").font = Font(name=FONT, bold=True)
    for j, t in enumerate(tags, start=r0 + 1):
        ws.cell(j, 1, t).font = Font(name=FONT)
        ws.cell(j, 2, TAG_KO.get(t, t)).font = Font(name=FONT)

    ws = wb.create_sheet("표시")
    head = ["번호", "상품id", "출처", "카테고리", "상품명", "이미지 1", "이미지 2", "이미지 3", "판매자 원문", "이미지에서 자동으로 읽은 글(오독 가능)"]
    n_fixed = len(head)
    head += [f"{t}\n{TAG_KO.get(t, '')}" for t in tags] + ["메모"]
    for c, h in enumerate(head, start=1):
        cell = ws.cell(1, c, h)
        cell.font = Font(name=FONT, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1='"1,0,?"', allow_blank=False, showErrorMessage=True,
                        errorTitle="입력 오류", error="1, 0, ? 중 하나만 입력하세요.")
    ws.add_data_validation(dv)
    for i, it in enumerate(rows, start=2):
        vals = [i - 1, it["id"], it["src"], it["cat"], it["name"]]
        imgs = it.get("images", [])
        vals += [imgs[k] if k < len(imgs) else "" for k in range(3)]
        vals += [(it.get("text") or "")[:4000], (it.get("vis") or "")[:4000]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(i, c, v)
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if 6 <= c <= 8 and v:
                cell.hyperlink = v
                cell.font = Font(name=FONT, color="0563C1", underline="single")
        for c in range(n_fixed + 1, n_fixed + len(tags) + 1):
            cell = ws.cell(i, c)
            cell.fill = yellow
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(horizontal="center", vertical="top")
            dv.add(cell)
        ws.row_dimensions[i].height = 150
    widths = [6, 11, 8, 10, 38, 22, 22, 22, 55, 55] + [9] * len(tags) + [30]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 60
    ws.freeze_panes = ws.cell(2, 6)
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=24)
    ap.add_argument("--seed", type=int, default=20261012)
    args = ap.parse_args()
    items = json.loads((ROOT / "samples" / "real_50.json").read_text(encoding="utf-8"))["items"]
    vis = {x["id"]: x.get("text", "") for x in json.loads((ROOT / "samples" / "real_50_vision.json").read_text(encoding="utf-8"))["items"]}
    for it in items:
        it["vis"] = vis.get(it["id"], "")
    rows = pick(items, args.n, args.seed)
    rules = matching.load_rules()
    tags = sorted(rules.weights, key=lambda t: (t.split(":")[0] != "fit", t))
    outdir = ROOT / "samples" / "labeling"
    outdir.mkdir(exist_ok=True)
    build(outdir / "label_A.xlsx", rows, tags, "표시자 A")
    build(outdir / "label_B.xlsx", rows, tags, "표시자 B")
    (outdir / "items_manifest.json").write_text(json.dumps({"seed": args.seed, "ids": [r["id"] for r in rows], "tags": tags}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(rows)}개 상품, 태그 {len(tags)}개 -> {outdir}")
    from collections import Counter
    print(dict(Counter((r["src"], r["cat"]) for r in rows)))


if __name__ == "__main__":
    main()
